from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def _lines(text: str) -> list[str]:
    return [line.strip() for line in text.splitlines()]


def _is_heading(line: str) -> bool:
    if not line:
        return False
    if len(line) <= 80 and line.upper() == line and re.search(r"[A-Z]", line):
        return True
    return bool(re.match(r"^\d+[\.)]\s+[A-Z]", line))


def format_docx(text: str, doc_type: str, logo_bytes: bytes | None = None) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    if logo_bytes:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(BytesIO(logo_bytes), width=Inches(1.2))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for line in _lines(text):
        if not line:
            doc.add_paragraph()
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        if _is_heading(line):
            r = p.add_run(line)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
        else:
            r = p.add_run(line)
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase — AI-generated draft. Review with a qualified legal professional before use.").italic = True

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def format_pdf(text: str, doc_type: str, logo_bytes: bytes | None = None) -> bytes:
    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=22 * mm,
        bottomMargin=20 * mm,
        title=doc_type,
        author="LegalEase",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "LegalTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=19,
        alignment=TA_CENTER,
        spaceAfter=14,
    )
    heading_style = ParagraphStyle(
        "LegalHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        spaceBefore=7,
        spaceAfter=5,
    )
    body_style = ParagraphStyle(
        "LegalBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=14,
        spaceAfter=6,
    )

    story = []
    if logo_bytes:
        from reportlab.platypus import Image
        from PIL import Image as PILImage
        image = PILImage.open(BytesIO(logo_bytes))
        w, h = image.size
        max_w, max_h = 35 * mm, 25 * mm
        scale = min(max_w / w, max_h / h)
        story.append(Image(BytesIO(logo_bytes), width=w * scale, height=h * scale))
        story.append(Spacer(1, 4))

    story.append(Paragraph(_escape(doc_type.upper()), title_style))

    for line in _lines(text):
        if not line:
            story.append(Spacer(1, 4))
        elif _is_heading(line):
            story.append(Paragraph(_escape(line), heading_style))
        else:
            story.append(Paragraph(_escape(line), body_style))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.drawCentredString(
            A4[0] / 2,
            10 * mm,
            "LegalEase — AI-generated draft | Review before use | Page %d" % document.page,
        )
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _escape(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )
